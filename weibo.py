"""
Weibo HTML → Free Google Translate → Excel
=============================================
Reads Weibo hot-search HTML files saved by SingleFile,
calls a free Google Translate interface to translate Chinese → English,
then appends data to a master Excel workbook with standard clean formatting.

Requirements:
    pip install beautifulsoup4 openpyxl pandas deep-translator
"""

import argparse
import glob
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup, Comment
from deep_translator import GoogleTranslator
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

# ── Border style ──────────────────────────────────────────────
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ════════════════════════════════════════════════════════════════
# 1. HTML PARSING
# ════════════════════════════════════════════════════════════════

def extract_date_from_html(soup) -> datetime | None:
    for c in soup.find_all(string=lambda t: isinstance(t, Comment)):
        m = re.search(r"saved date:\s*(.+?)(?:\n|$)", str(c))
        if m:
            raw = m.group(1).strip()
            raw = re.sub(r"\s+GMT[^\s]*.*$", "", raw)
            try:
                return datetime.strptime(raw, "%a %b %d %Y %H:%M:%S")
            except ValueError:
                pass
    return None


def _parse_rank(td01) -> int | None:
    if not td01:
        return None
    text = td01.get_text(strip=True)
    if text.isdigit():
        return int(text)
    return None


def _parse_engagement(td02) -> int | None:
    """
    提取纯数字热度值。
    - 如果找到数字热度，返回整数 (int)。
    - 如果该行没有数字热度值，一律返回 None（在 Excel 中保持空白）。
    """
    if not td02:
        return None

    # 提取包含数字的 span 标签或检查文本
    span_tag = td02.select_one("span")
    if span_tag:
        span_text = span_tag.get_text(strip=True).replace(",", "")
        num_match = re.search(r'\d+', span_text)
        if num_match:
            return int(num_match.group())

    return None


def parse_html_file(path: str) -> dict:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    date_val = extract_date_from_html(soup)
    if date_val is None:
        stem = Path(path).stem
        try:
            date_val = datetime.strptime(stem, "%Y-%m-%d_%H%M")
        except ValueError:
            date_val = datetime.now()

    entries = []

    for row in soup.select("tbody tr"):
        td01 = row.select_one(".td-01")
        td02 = row.select_one(".td-02")
        if not td02:
            continue

        a_tag = td02.select_one("a")
        if not a_tag:
            continue

        chinese = a_tag.get_text(strip=True)
        if not chinese or chinese == "查看更多":
            continue

        rank = _parse_rank(td01)
        engagement = _parse_engagement(td02)

        entries.append({
            "rank":       rank,
            "chinese":    chinese,
            "engagement": engagement,
        })

    return {"date": date_val, "entries": entries, "source": Path(path).name}


# ════════════════════════════════════════════════════════════════
# 2. TOPIC TRANSLATION
# ════════════════════════════════════════════════════════════════

def translate_entries(entries: list[dict]) -> list[dict]:
    enriched = []
    translator = GoogleTranslator(source='zh-CN', target='en')

    total = len(entries)
    print(f"  Starting translation for {total} topics…")

    for i, entry in enumerate(entries, 1):
        out = dict(entry)
        chinese_text = entry["chinese"]
        clean_text = chinese_text.strip('#')

        try:
            translated = translator.translate(clean_text)
            out["english"] = translated.title() if translated else chinese_text
            print(f"    [{i}/{total}] {chinese_text} → {out['english']}")
        except Exception as e:
            print(f"    [{i}/{total}] Translation failed for '{chinese_text}': {e}")
            out["english"] = chinese_text

        enriched.append(out)

        if i < total:
            time.sleep(0.3)

    return enriched


# ════════════════════════════════════════════════════════════════
# 3. EXCEL OUTPUT
# ════════════════════════════════════════════════════════════════

def _style_standard(ws, row, col, value=None, bold=False, align="left", wrap=False, fmt=None):
    c = ws.cell(row=row, column=col)
    if value is not None:
        c.value = value
    c.font = Font(bold=bold, name="Arial", size=10)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = BORDER
    if fmt:
        c.number_format = fmt
    return c


def append_to_excel(html_data_list, output_path):
    file_path = Path(output_path)

    if file_path.exists():
        print(f"  Appending to existing file: {output_path}")
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active

        real_max_row = 1
        for row in range(ws.max_row, 0, -1):
            if any(ws.cell(row=row, column=col).value is not None for col in range(1, 6)):
                real_max_row = row
                break
        is_empty_file = False
    else:
        print(f"  Creating new workbook: {output_path}")
        wb = Workbook()
        ws = wb.active
        ws.title = "All Data"
        real_max_row = 1
        is_empty_file = True

    if is_empty_file or real_max_row == 1:
        headers = [
            "Date",
            "Rank",
            "Topic Headline: Chinese",
            "Topic Headline: English Translation",
            "Engagement",
        ]
        for c, label in enumerate(headers, 1):
            _style_standard(ws, 1, c, label, bold=True, align="center", wrap=True)
        ws.row_dimensions[1].height = 28
        real_max_row = 1

    for index, file_data in enumerate(html_data_list):

        if (real_max_row > 1 and index == 0) or (index > 0):
            empty_row = real_max_row + 1
            for c in range(1, 6):
                _style_standard(ws, empty_row, c, value="")
            ws.row_dimensions[empty_row].height = 65
            start_row = empty_row + 1
        else:
            start_row = real_max_row + 1

        for r_offset, row in enumerate(file_data):
            er = start_row + r_offset

            _style_standard(ws, er, 1, row["date"], align="center", fmt="YYYY-MM-DD HH:MM")

            rank_val = row.get("rank")
            _style_standard(ws, er, 2, rank_val, align="center")

            _style_standard(ws, er, 3, row["chinese"], wrap=True)
            _style_standard(ws, er, 4, row["english"], wrap=True)

            eng = row.get("engagement")

            if isinstance(eng, int):
                # 纯热度数字，右对齐加千分位
                _style_standard(ws, er, 5, eng, align="right", fmt="#,##0")
            else:
                # 没有任何数字的（无论是空白行还是文本标签行），直接留空单元格
                _style_standard(ws, er, 5, None, align="right")

        real_max_row = start_row + len(file_data) - 1

    widths = {1: 18, 2: 8, 3: 35, 4: 50, 5: 16}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = "A2"
    wb.save(output_path)


# ════════════════════════════════════════════════════════════════
# 4. MAIN
# ════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="Weibo HTML → Free Google Translate → Excel"
    )
    parser.add_argument("--input",  "-i", required=True,
                        help="HTML file or folder containing .html files")
    parser.add_argument("--output", "-o", default="weibo_master.xlsx",
                        help="Output Excel file (appended if it already exists)")
    args = parser.parse_args()

    p = Path(args.input)
    if p.is_file():
        html_files = [str(p)]
    elif p.is_dir():
        html_files = sorted(
            glob.glob(str(p / "*.html")) + glob.glob(str(p / "*.htm"))
        )
        if not html_files:
            sys.exit(f"No HTML files found in {args.input}")
    else:
        sys.exit(f"Path not found: {args.input}")

    print(f"\n  Found {len(html_files)} HTML file(s)\n")

    html_data_list = []

    for path in html_files:
        print(f"  Parsing: {Path(path).name}")
        parsed = parse_html_file(path)
        entries = parsed["entries"]
        print(f"  → {len(entries)} entries, date={parsed['date']}")

        if not entries:
            continue

        translated = translate_entries(entries)

        for entry in translated:
            entry["date"] = parsed["date"]

        html_data_list.append(translated)
        print()

    if not html_data_list:
        sys.exit("No new data extracted.")

    total_rows = sum(len(x) for x in html_data_list)
    print(f"  Writing {total_rows} total entries to Excel…")
    append_to_excel(html_data_list, args.output)
    print(f"\n  Done: {args.output}")


if __name__ == "__main__":
    main()